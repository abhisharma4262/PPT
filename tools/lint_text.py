"""
lint_text: checks every slide for the things the deck owner has had to fix by hand.

  python3 lint_text.py deck.pptx        (or deck.html)
  python3 lint_text.py deck.pptx --no-stats    also flags percentages and statistics

ERRORS (must fix): em dash, en dash used as punctuation, double hyphen, emoji, exclamation mark,
                   dots or bullets used to separate items in a line ("A · B · C").
WARNINGS (review): AI-sounding words and constructions, arrogant phrasing, long titles, text below 9 pt.
Exit code 1 when any error is found.
"""
import re, sys, html as htmlmod

ERR = [
    (r'\u2014', 'em dash: rewrite with ";" "," ":" or two sentences'),
    (r'\s[\u00b7\u2022]\s', 'dot or bullet used as a separator: use separate items (mkit band_items) or rewrite'),
    (r'\s\u2013\s', 'en dash used as a dash: rewrite with ";" or ","'),
    (r'(?<!-)--(?!-)', 'double hyphen used as a dash'),
    (r'[\U0001F300-\U0001FAFF☀-➿]', 'emoji'),
    (r'!', 'exclamation mark'),
]
WARN = [
    (r'\b(delve|unleash|game[- ]changer|cutting[- ]edge|revolutioni[sz]e|transformative|tapestry|elevate|'
     r'empower|synerg\w*|paradigm|best[- ]in[- ]class|world[- ]class|next[- ]gen\w*|state[- ]of[- ]the[- ]art|'
     r'supercharge|turbocharge|robust|effortless(ly)?|ever[- ]evolving|fast[- ]paced)\b',
     'AI-sounding or hype word'),
    (r"\b(in today's|harness the power|leverage the power|at the end of the day|it's not just|not just .{1,40}? but|"
     r"not only .{1,40}? but also|moreover|furthermore|in conclusion|needless to say)\b", 'AI-sounding construction'),
    (r'\b(genuinely|honestly|truly|really|very)\b', 'filler intensifier'),
    (r'\b(one-off project|only we|nobody else|unmatched|unrivalled|unparalleled|guarantee[ds]?)\b', 'arrogant or absolute claim'),
]
STATS = (r'\b\d{1,3}(\.\d+)?\s?%|\b\d+x\b', 'statistic or percentage (allowed only with a cited source and the user\'s OK)')

def slides_from_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                yield from walk(sh.shapes)
            else:
                yield sh
    for s in prs.slides:
        texts = []
        title = ''
        small = []
        for sh in walk(s.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text
                if '©' in t:
                    continue
                texts.append(t)
                if not sh.is_placeholder:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.size and r.font.size.pt < 9 and r.text.strip():
                                small.append(f'{r.font.size.pt:g} pt: {r.text.strip()[:40]}')
                sz = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                if sz and max(sz) >= 24 and (sh.top or 0) < 914400:
                    title = t
        out.append((title, texts, small))
    return out

def slides_from_html(path):
    h = open(path, encoding='utf-8').read()
    h = re.sub(r'<style.*?</style>|<script.*?</script>', '', h, flags=re.S)
    out = []
    for sec in re.findall(r'<section class="frame".*?</section>', h, flags=re.S):
        sec = re.sub(r'<div class="tpl[^"]*">.*?</div>', '', sec, flags=re.S)
        m = re.search(r'<h1 class="title[^"]*">(.*?)</h1>', sec, flags=re.S)
        title = htmlmod.unescape(re.sub(r'<[^>]+>', '', m.group(1))) if m else ''
        txt = htmlmod.unescape(re.sub(r'<[^>]+>', '\n', sec))
        out.append((title, [t for t in txt.split('\n') if t.strip()], []))
    return out

def main():
    path = sys.argv[1]
    no_stats = '--no-stats' in sys.argv
    slides = slides_from_pptx(path) if path.lower().endswith('.pptx') else slides_from_html(path)
    errors = 0
    for i, (title, texts, small) in enumerate(slides, 1):
        found = []
        for sm in small[:6]:
            found.append(('warn', 'text below the 9 pt floor', sm))
        for t in texts:
            for pat, msg in ERR:
                for m in re.finditer(pat, t):
                    found.append(('ERROR', msg, t[max(0, m.start()-30):m.end()+30])); errors += 1
            for pat, msg in WARN:
                for m in re.finditer(pat, t, flags=re.I):
                    found.append(('warn', msg, t[max(0, m.start()-30):m.end()+30]))
            if no_stats:
                for m in re.finditer(STATS[0], t):
                    found.append(('warn', STATS[1], t[max(0, m.start()-30):m.end()+30]))
        if title and len(title.replace('\n', ' ')) > 68:
            found.append(('warn', f'title is {len(title)} characters; one line holds about 65 at 28pt bold', title))
        for lvl, msg, ctx in found:
            print(f'slide {i:>2} {lvl:5} {msg}: "...{ctx.strip()}..."')
    print(f'{len(slides)} slides checked, {errors} errors')
    sys.exit(1 if errors else 0)

if __name__ == '__main__':
    main()
