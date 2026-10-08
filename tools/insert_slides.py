"""
insert_slides: put newly built slides into a deck the user has already edited, without rebuilding it.

Build the new slides with mkit and html2pptx as usual (raw.pptx), then:

  # add raw slides 1-3 right after slide 8 of the user's deck (they take slide 8's layout)
  python3 insert_slides.py deck.pptx raw.pptx out.pptx --after 8 --take 1,2,3

  # replace everything on slide 5 with raw slide 1 (keeps slide 5's layout, logo and page number)
  python3 insert_slides.py deck.pptx raw.pptx out.pptx --replace 5:1

  # replace only a region of slide 3 (shapes whose top is between y0 and y1, canvas px 0..1080)
  # with the shapes of raw slide 2; build raw slide 2 with only that region's content
  python3 insert_slides.py deck.pptx raw.pptx out.pptx --region 3:2:820:1080

  # delete slides (after any inserts, numbers refer to the final deck)
  python3 insert_slides.py deck.pptx raw.pptx out.pptx --delete 9,10

Options can be combined; they run in the order: region, replace, insert, delete.
Slide numbers are 1-based. Prints the resulting slide list (eyebrow and title) for a quick check.
"""
import argparse, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from merge_into_template import clone_slide, import_shapes, renumber_ids, move_to, drop_slide

PX = 914400 / 144   # EMU per canvas px


def clear(slide, y0=None, y1=None):
    for sh in list(slide.shapes):
        if sh.is_placeholder:
            continue
        if y0 is not None and not (y0 <= (sh.top or 0) / PX < y1):
            continue
        sh._element.getparent().remove(sh._element)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('deck'); ap.add_argument('raw'); ap.add_argument('out')
    ap.add_argument('--after', type=int, help='insert the --take slides after this deck slide')
    ap.add_argument('--take', default='', help='raw slide numbers to insert, e.g. 1,2,3')
    ap.add_argument('--replace', action='append', default=[], help='N:k  replace deck slide N content with raw slide k')
    ap.add_argument('--region', action='append', default=[], help='N:k:y0:y1  replace shapes on deck slide N whose top is in [y0,y1) px')
    ap.add_argument('--delete', default='', help='deck slide numbers to delete at the end')
    a = ap.parse_args()

    prs = Presentation(a.deck)
    raw = Presentation(a.raw)
    pkg = prs.part.package
    slides = list(prs.slides)
    rs = list(raw.slides)

    for r in a.region:
        n, k, y0, y1 = r.split(':')
        s = slides[int(n) - 1]
        clear(s, float(y0), float(y1))
        import_shapes(s, rs[int(k) - 1], pkg)
        renumber_ids(s)

    for r in a.replace:
        n, k = (int(v) for v in r.split(':'))
        s = slides[n - 1]
        clear(s)
        import_shapes(s, rs[k - 1], pkg)
        renumber_ids(s)

    if a.after and a.take:
        base = slides[a.after - 1]
        for i, k in enumerate(int(v) for v in a.take.split(',')):
            ns = clone_slide(prs, base)
            clear(ns)
            import_shapes(ns, rs[k - 1], pkg)
            renumber_ids(ns)
            move_to(prs, ns, a.after + i)

    if a.delete:
        cur = list(prs.slides)
        for n in sorted((int(v) for v in a.delete.split(',')), reverse=True):
            drop_slide(prs, cur[n - 1])

    prs.save(a.out)
    for i, s in enumerate(prs.slides, 1):
        t = [sh.text_frame.text.replace('\x0b', ' ') for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()][:2]
        print(f'{i:>2}  ' + '  |  '.join(x[:70] for x in t))
    print('written', a.out)


if __name__ == '__main__':
    main()
